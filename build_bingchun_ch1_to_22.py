from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "BingChun_Chuong1_den_2_2.docx"


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("left", "top", "right", "bottom", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def style_run(run, bold=False, size=13, color="000000", italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.rPr
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run._element.insert(0, rpr)
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")


def add_paragraph(doc, text="", style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_cm=0.75, before=0, after=6, line_spacing=1.5, bold=False, size=13):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(first_line_cm) if first_line_cm else Cm(0)
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        style_run(r, bold=bold, size=size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = p.paragraph_format
        fmt.space_before = Pt(12)
        fmt.space_after = Pt(6)
        fmt.line_spacing = 1.2
        fmt.first_line_indent = Cm(0)
        r = p.add_run(text)
        style_run(r, bold=True, size=15)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = p.paragraph_format
        fmt.space_before = Pt(8)
        fmt.space_after = Pt(4)
        fmt.line_spacing = 1.2
        fmt.first_line_indent = Cm(0)
        r = p.add_run(text)
        style_run(r, bold=True, size=13.5)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = p.paragraph_format
        fmt.space_before = Pt(6)
        fmt.space_after = Pt(4)
        fmt.line_spacing = 1.2
        fmt.first_line_indent = Cm(0)
        r = p.add_run(text)
        style_run(r, bold=True, size=13)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style=None)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75 + level * 0.6)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.45
    r = p.add_run("• " + text)
    style_run(r, size=13)
    return p


def add_numbered(doc, number, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.45
    r1 = p.add_run(f"{number}. ")
    style_run(r1, bold=True, size=13)
    r2 = p.add_run(text)
    style_run(r2, size=13)
    return p


doc = Document()

# Page setup
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(3.0)
section.right_margin = Cm(2.0)
section.header_distance = Cm(1.0)
section.footer_distance = Cm(1.0)

# Default font
styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"].font.size = Pt(13)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.5
styles["Normal"].paragraph_format.first_line_indent = Cm(0.75)

# Title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.space_before = Pt(0)
r = p.add_run("BÁO CÁO PHÂN TÍCH THIẾT KẾ HỆ THỐNG")
style_run(r, bold=True, size=18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
r = p.add_run("DỰ ÁN BINGCHUN")
style_run(r, bold=True, size=22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("TỪ CHƯƠNG 1 ĐẾN HẾT MỤC 2.2")
style_run(r, bold=True, size=15)

for _ in range(3):
    doc.add_paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sinh viên thực hiện")
style_run(r, size=13)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("NGUYỄN ĐỨC NGHĨA")
style_run(r, bold=True, size=16)

for _ in range(6):
    doc.add_paragraph("")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Năm 2026")
style_run(r, size=13)

doc.add_page_break()

# Header/footer
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer)

# Chapter 1
add_heading(doc, "CHƯƠNG 1  TÌM HIỂU BÀI TOÁN", 1)

add_heading(doc, "1.1. Giới thiệu bài toán", 2)
add_paragraph(
    doc,
    "Dự án BingChun được xây dựng nhằm số hóa quy trình bán hàng trực tuyến cho một cửa hàng đồ uống "
    "trà sữa quy mô nhỏ đến vừa, nơi khách hàng có nhu cầu đặt món nhanh, tùy biến sản phẩm linh hoạt "
    "và nhận xác nhận đơn hàng tức thời. Trong bối cảnh khách hàng ngày càng quen với việc đặt đồ uống "
    "trực tuyến qua website thay vì trao đổi thủ công trên tin nhắn, một hệ thống tập trung là cần thiết "
    "để chuẩn hóa dữ liệu sản phẩm, tối ưu xử lý đơn hàng và giảm sai sót trong quá trình phục vụ.",
)
add_paragraph(
    doc,
    "Trước khi có hệ thống, các thao tác như tra cứu menu, báo giá, ghi chú sở thích, xác nhận trạng thái "
    "đơn hàng hay phản hồi đánh giá thường dễ bị phân tán ở nhiều kênh khác nhau. Điều này làm tăng nguy cơ "
    "sai lệch thông tin, làm chậm tốc độ phục vụ và khiến việc quản trị trở nên khó kiểm soát khi số lượng "
    "đơn tăng lên. Từ thực tế đó, đề tài BingChun được lựa chọn với mục tiêu xây dựng một nền tảng web "
    "tích hợp đầy đủ các chức năng bán hàng, thanh toán, chăm sóc khách hàng và quản trị vận hành.",
)
add_numbered(
    doc,
    "Lý do chọn bài toán",
    "Bài toán được chọn vì có tính thực tiễn cao, phù hợp với xu hướng mua hàng trực tuyến của người dùng hiện nay "
    "và phản ánh đúng nhu cầu vận hành của một cửa hàng đồ uống. Hệ thống không chỉ phục vụ thao tác đặt món của "
    "khách hàng mà còn hỗ trợ nhân viên quản lý danh mục sản phẩm, đơn hàng, voucher, banner và báo cáo doanh thu. "
    "Đây là một bài toán vừa có tính nghiệp vụ rõ ràng, vừa đủ phạm vi để triển khai thành sản phẩm hoàn chỉnh trong "
    "đồ án học phần.",
)
add_numbered(
    doc,
    "Đề xuất các giải pháp thực hiện",
    "Giải pháp đề xuất là xây dựng một hệ thống web tách biệt giữa giao diện người dùng và tầng xử lý nghiệp vụ. "
    "Phía người dùng cho phép duyệt menu, tìm kiếm sản phẩm theo thời gian thực, lọc theo danh mục, mức giá và đánh giá, "
    "xem chi tiết sản phẩm, tùy chỉnh size, topping, ghi chú, quản lý giỏ hàng và đặt hàng trực tiếp. Phía quản trị cung "
    "cấp chức năng quản lý sản phẩm, loại sản phẩm, banner, mã giảm giá, đơn hàng, đánh giá, báo cáo và nhật ký hệ thống. "
    "Các thông báo quan trọng được gửi qua email và WebSocket để đảm bảo người dùng lẫn quản trị viên nhận được phản hồi "
    "kịp thời ngay khi phát sinh sự kiện nghiệp vụ.",
)
add_numbered(
    doc,
    "Đánh giá tính khả thi của hệ thống",
    "Về mặt kỹ thuật, hệ thống có tính khả thi cao vì các yêu cầu nghiệp vụ đều có thể triển khai bằng kiến trúc web phổ biến "
    "với cơ sở dữ liệu quan hệ, API phía máy chủ và giao diện người dùng tách riêng. Các chức năng như đăng nhập, giỏ hàng, "
    "đặt hàng, thanh toán, gửi email, thông báo thời gian thực và quản trị dữ liệu đều là những nghiệp vụ đã có mẫu triển khai "
    "rõ ràng trong các dự án thương mại điện tử và đặt hàng trực tuyến. Về mặt vận hành, phạm vi của dự án phù hợp với quy mô "
    "một cửa hàng đồ uống, không đòi hỏi hạ tầng quá lớn nhưng vẫn đủ để mô phỏng quy trình nghiệp vụ hoàn chỉnh. Về mặt kinh tế, "
    "giải pháp có chi phí triển khai thấp vì tận dụng công nghệ web mã nguồn phổ biến và có thể vận hành trên hạ tầng cơ bản.",
)
add_numbered(
    doc,
    "Lập kế hoạch thực hiện",
    "Kế hoạch triển khai dự án được chia thành các giai đoạn rõ ràng để đảm bảo tiến độ và chất lượng. Giai đoạn đầu tập trung "
    "khảo sát yêu cầu, phân tích nghiệp vụ và xác định phạm vi. Giai đoạn tiếp theo là thiết kế cơ sở dữ liệu, luồng nghiệp vụ "
    "và giao diện. Sau đó tiến hành phát triển từng phân hệ chính, kiểm thử chức năng và kiểm thử tích hợp. Giai đoạn cuối là "
    "hoàn thiện báo cáo, rà soát lỗi giao diện, kiểm tra dữ liệu và chuẩn bị bản trình bày. Cách tổ chức này giúp kiểm soát công việc "
    "theo từng mốc, đồng thời dễ dàng điều chỉnh khi có thay đổi về yêu cầu.",
)

add_heading(doc, "1.2. Tìm hiểu yêu cầu người dùng", 2)
add_paragraph(
    doc,
    "Việc tìm hiểu yêu cầu người dùng là cơ sở quan trọng để hệ thống BingChun bám sát nhu cầu thực tế. Nhóm thực hiện khảo sát "
    "theo hướng kết hợp giữa quan sát quy trình đặt hàng hiện tại, phân tích các yêu cầu lặp lại trong quá trình bán hàng và xác "
    "định các điểm gây khó khăn cho người dùng. Từ đó, danh sách yêu cầu được chuẩn hóa thành các nhóm chức năng phục vụ khách "
    "hàng, nhân viên và quản trị viên.",
)
add_heading(doc, "1.2.1. Kế hoạch xác định yêu cầu người dùng", 3)
add_paragraph(
    doc,
    "Kế hoạch xác định yêu cầu được thực hiện theo ba bước. Bước thứ nhất là thu thập thông tin ban đầu từ các thao tác bán hàng "
    "thường gặp như xem menu, chọn món, tính tiền, xác nhận đơn và phản hồi sau bán hàng. Bước thứ hai là mô hình hóa quy trình "
    "nghiệp vụ để xác định điểm vào, điểm ra và trạng thái chuyển tiếp của từng tác vụ. Bước thứ ba là tổng hợp yêu cầu thành "
    "những chức năng cụ thể, phân loại theo nhóm người dùng và mức độ ưu tiên triển khai.",
)
add_heading(doc, "1.2.2. Quy trình nghiệp vụ hiện tại", 3)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "Bước"
hdr[1].text = "Nội dung thực hiện"
hdr[2].text = "Người thực hiện"
hdr[3].text = "Hồ sơ, biểu mẫu liên quan"
set_repeat_table_header(table.rows[0])
for c in hdr:
    shade_cell(c, "D9E2F3")
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in c.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            style_run(r, bold=True, size=11.5)
    set_cell_border(c, left={"val": "single", "sz": 6, "color": "D9D9D9"},
                    top={"val": "single", "sz": 6, "color": "D9D9D9"},
                    right={"val": "single", "sz": 6, "color": "D9D9D9"},
                    bottom={"val": "single", "sz": 6, "color": "D9D9D9"})

rows = [
    ("1", "Khách hàng truy cập website, xem menu, tìm kiếm sản phẩm và lọc theo danh mục, giá hoặc mức đánh giá.", "Khách hàng", "Trang danh sách sản phẩm, dữ liệu danh mục, bộ lọc"),
    ("2", "Khách hàng chọn sản phẩm, tùy chỉnh size, topping, ghi chú mua hàng và thêm vào giỏ hàng.", "Khách hàng", "Chi tiết sản phẩm, giỏ hàng tạm thời"),
    ("3", "Khách hàng kiểm tra lại đơn, chọn địa chỉ nhận hàng, voucher và phương thức thanh toán.", "Khách hàng", "Giỏ hàng, địa chỉ giao hàng, phiếu giảm giá"),
    ("4", "Hệ thống tạo đơn hàng, ghi nhận thanh toán, trừ số lượng tồn kho và gửi thông báo xác nhận.", "Hệ thống, quản trị viên", "Đơn hàng, hóa đơn, lịch sử trạng thái, email xác nhận"),
    ("5", "Sau khi đơn hoàn tất, khách hàng đánh giá sản phẩm và có thể xem hoặc chỉnh sửa thông tin tài khoản.", "Khách hàng, quản trị viên", "Đánh giá sản phẩm, hồ sơ khách hàng, điểm tích lũy"),
]
for row_data in rows:
    row = table.add_row().cells
    for i, val in enumerate(row_data):
        row[i].text = val
        row[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in row[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                style_run(r, size=11)
        set_cell_border(row[i], left={"val": "single", "sz": 6, "color": "D9D9D9"},
                        top={"val": "single", "sz": 6, "color": "D9D9D9"},
                        right={"val": "single", "sz": 6, "color": "D9D9D9"},
                        bottom={"val": "single", "sz": 6, "color": "D9D9D9"})

doc.add_paragraph("")

add_heading(doc, "1.2.3. Đánh giá quy trình hiện tại và đề xuất cải tiến", 3)
add_paragraph(
    doc,
    "Quy trình đặt hàng thủ công bằng tin nhắn hoặc ghi chép rời rạc có ưu điểm là dễ bắt đầu nhưng bộc lộ nhiều hạn chế khi "
    "khách hàng tăng lên. Việc kiểm tra sản phẩm, xác nhận size, topping và địa chỉ theo cách thủ công dễ phát sinh nhầm lẫn, "
    "đặc biệt vào thời điểm cao điểm. Ngoài ra, khách hàng khó biết trạng thái đơn của mình đang ở bước nào, trong khi nhân viên "
    "phải đối chiếu nhiều nguồn thông tin khác nhau để xử lý một đơn hàng hoàn chỉnh.",
)
add_paragraph(
    doc,
    "Từ các hạn chế đó, hệ thống BingChun được đề xuất cải tiến theo hướng tập trung hóa dữ liệu và tự động hóa các bước lặp lại. "
    "Khách hàng có thể tự thao tác trên giao diện web, hệ thống tự tính tiền, lưu lịch sử mua hàng, cập nhật số lượng sản phẩm, "
    "gửi email xác nhận và phát thông báo theo thời gian thực. Quản trị viên có thể theo dõi đơn hàng, phản hồi đánh giá, quản lý "
    "sản phẩm và báo cáo mà không cần đối chiếu thủ công trên nhiều kênh khác nhau. Cách tiếp cận này giúp giảm lỗi nghiệp vụ, "
    "tăng tốc độ phục vụ và nâng cao trải nghiệm người dùng.",
)

# Chapter 2
add_heading(doc, "CHƯƠNG 2  PHÂN TÍCH HỆ THỐNG VÀ THIẾT KẾ HỆ THỐNG", 1)

add_heading(doc, "2.1. Xác định yêu cầu chức năng và phi chức năng của hệ thống", 2)
add_paragraph(
    doc,
    "Trên cơ sở khảo sát nghiệp vụ thực tế, hệ thống BingChun được xác định với hai nhóm yêu cầu chính là yêu cầu chức năng và "
    "yêu cầu phi chức năng. Yêu cầu chức năng mô tả các việc hệ thống phải thực hiện để phục vụ khách hàng và quản trị viên; "
    "yêu cầu phi chức năng mô tả những tiêu chí chất lượng mà hệ thống phải đạt được trong quá trình vận hành.",
)

add_heading(doc, "2.1.1. Yêu cầu chức năng", 3)
add_bullet(doc, "Khách vãng lai có thể xem danh sách sản phẩm, xem chi tiết từng món, tra cứu theo danh mục và tìm kiếm theo từ khóa.")
add_bullet(doc, "Khách hàng đăng ký và đăng nhập tài khoản để quản lý thông tin cá nhân, địa chỉ nhận hàng và lịch sử mua hàng.")
add_bullet(doc, "Khách hàng thêm sản phẩm vào giỏ hàng, thay đổi số lượng, chọn size, topping, ghi chú riêng và áp dụng mã giảm giá.")
add_bullet(doc, "Hệ thống tính toán tổng tiền, phí giao hàng, điểm tích lũy, số tiền được giảm và số tiền thanh toán cuối cùng.")
add_bullet(doc, "Hệ thống hỗ trợ hai phương thức thanh toán chính gồm thanh toán khi nhận hàng và thanh toán trực tuyến theo cổng tích hợp.")
add_bullet(doc, "Sau khi đơn hàng được xác nhận, hệ thống gửi email và thông báo realtime cho khách hàng để cập nhật trạng thái.")
add_bullet(doc, "Khách hàng có thể đánh giá sản phẩm sau khi hoàn thành đơn và xem lại nội dung đã đánh giá trước đó.")
add_bullet(doc, "Quản trị viên quản lý sản phẩm, danh mục, banner, voucher, đơn hàng, phản hồi đánh giá, báo cáo doanh thu và nhật ký hoạt động.")

add_heading(doc, "2.1.2. Yêu cầu phi chức năng", 3)
add_bullet(doc, "Hệ thống phải phản hồi nhanh, bảo đảm thao tác tìm kiếm, lọc và tải danh sách sản phẩm trong thời gian chấp nhận được.")
add_bullet(doc, "Dữ liệu người dùng, đơn hàng và thanh toán phải được lưu trữ an toàn, có cơ chế kiểm tra đầu vào và xác thực người dùng.")
add_bullet(doc, "Giao diện phải dễ hiểu, tương thích với nhiều kích thước màn hình và thuận tiện cho người dùng phổ thông.")
add_bullet(doc, "Kiến trúc hệ thống cần dễ bảo trì, cho phép bổ sung chức năng mới mà không làm ảnh hưởng lớn đến các phần hiện có.")
add_bullet(doc, "Các thông báo nghiệp vụ quan trọng như đặt hàng, xác nhận đơn, phản hồi đánh giá phải được truyền tải ổn định.")
add_bullet(doc, "Hệ thống cần có khả năng mở rộng khi số lượng sản phẩm, đơn hàng và người dùng tăng lên trong tương lai.")

add_heading(doc, "2.2. Bản tuyên bố phạm vi", 2)
add_paragraph(
    doc,
    "Bản tuyên bố phạm vi của dự án BingChun xác định rõ giới hạn triển khai, các chức năng cốt lõi cần hoàn thành và các nội dung "
    "không thuộc phạm vi của đề tài. Việc xác định phạm vi ngay từ đầu giúp hệ thống giữ được trọng tâm nghiệp vụ, tránh mở rộng lan man "
    "và tạo cơ sở cho việc đánh giá mức độ hoàn thành của đồ án.",
)

add_heading(doc, "2.2.1. Mục tiêu phạm vi", 3)
add_paragraph(
    doc,
    "Phạm vi của dự án tập trung vào việc xây dựng một hệ thống web bán hàng đồ uống có khả năng hỗ trợ đầy đủ quy trình từ xem sản phẩm, "
    "chọn món, đặt hàng, thanh toán, thông báo xác nhận đến quản trị vận hành. Hệ thống được thiết kế để phục vụ đồng thời hai nhóm người "
    "dùng chính là khách hàng và quản trị viên, trong đó khách hàng thao tác trên giao diện bán hàng còn quản trị viên kiểm soát dữ liệu và "
    "theo dõi hoạt động kinh doanh.",
)

add_heading(doc, "2.2.2. Phạm vi chức năng của hệ thống", 3)
scope = doc.add_table(rows=1, cols=3)
scope.style = "Table Grid"
scope.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = scope.rows[0].cells
hdr[0].text = "Nhóm chức năng"
hdr[1].text = "Thuộc phạm vi"
hdr[2].text = "Mô tả ngắn"
set_repeat_table_header(scope.rows[0])
for c in hdr:
    shade_cell(c, "44546A")
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in c.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            style_run(r, bold=True, size=11, color="FFFFFF")
    set_cell_border(c, left={"val": "single", "sz": 6, "color": "D9D9D9"},
                    top={"val": "single", "sz": 6, "color": "D9D9D9"},
                    right={"val": "single", "sz": 6, "color": "D9D9D9"},
                    bottom={"val": "single", "sz": 6, "color": "D9D9D9"})

scope_rows = [
    ("Bán hàng phía khách", "Có", "Xem menu, tìm kiếm, lọc sản phẩm, xem chi tiết, thêm vào giỏ và đặt hàng."),
    ("Tùy biến sản phẩm", "Có", "Chọn size, topping, ghi chú và tính lại giá theo cấu hình đã chọn."),
    ("Quản lý tài khoản", "Có", "Đăng nhập, đăng ký, cập nhật hồ sơ, quản lý địa chỉ và điểm tích lũy."),
    ("Thanh toán và thông báo", "Có", "Xử lý COD, thanh toán trực tuyến, gửi email xác nhận và thông báo realtime."),
    ("Đánh giá sản phẩm", "Có", "Gửi đánh giá, lọc theo sao, xem phản hồi từ quản trị viên."),
    ("Quản trị hệ thống", "Có", "Quản lý sản phẩm, danh mục, voucher, banner, đơn hàng, báo cáo và audit log."),
    ("Ứng dụng di động riêng", "Không", "Dự án chỉ triển khai phiên bản web, chưa phát triển app mobile độc lập."),
    ("Tích hợp vận chuyển tự động", "Không", "Chưa tích hợp sâu với đơn vị giao hàng bên ngoài ở mức đồng bộ thời gian thực."),
    ("Phân hệ kho nâng cao", "Không", "Chưa xây dựng mô hình kho đa chi nhánh hoặc tối ưu tồn kho theo lô hàng."),
]
for i, row_data in enumerate(scope_rows):
    row = scope.add_row().cells
    for j, val in enumerate(row_data):
        row[j].text = val
        row[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in row[j].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                style_run(r, size=11)
        set_cell_border(row[j], left={"val": "single", "sz": 6, "color": "D9D9D9"},
                        top={"val": "single", "sz": 6, "color": "D9D9D9"},
                        right={"val": "single", "sz": 6, "color": "D9D9D9"},
                        bottom={"val": "single", "sz": 6, "color": "D9D9D9"})
        if i % 2 == 1:
            shade_cell(row[j], "F7F9FC")

add_heading(doc, "2.2.3. Sản phẩm chuyển giao", 3)
add_paragraph(
    doc,
    "Sản phẩm chuyển giao của dự án gồm bộ mã nguồn phía giao diện người dùng, bộ mã nguồn phía máy chủ, lược đồ cơ sở dữ liệu, "
    "bộ tài liệu phân tích thiết kế và bản báo cáo mô tả nghiệp vụ. Ngoài ra, hệ thống còn tạo ra dữ liệu đầu ra phục vụ vận hành "
    "như tài khoản người dùng, danh sách sản phẩm, đơn hàng, đánh giá, báo cáo doanh thu và nhật ký hoạt động.",
)
add_paragraph(
    doc,
    "Về mặt sử dụng, kết quả cuối cùng cần đáp ứng được việc khách hàng có thể đặt món trọn vẹn trên website, còn quản trị viên có thể "
    "theo dõi đơn hàng, sửa dữ liệu và nhận thông báo ngay khi có sự kiện nghiệp vụ mới phát sinh.",
)

add_heading(doc, "2.2.4. Giả định và ràng buộc", 3)
add_bullet(doc, "Người dùng có thiết bị kết nối internet và sử dụng trình duyệt web hiện đại.")
add_bullet(doc, "Dữ liệu sản phẩm, giá bán, topping và voucher được quản trị viên cập nhật trước khi vận hành chính thức.")
add_bullet(doc, "Hệ thống email và dịch vụ thông báo được cấu hình sẵn để đảm bảo gửi xác nhận đơn hàng và phản hồi trạng thái.")
add_bullet(doc, "Quy mô dự án phù hợp với mô hình cửa hàng đơn lẻ hoặc chuỗi cửa hàng nhỏ, chưa tối ưu cho vận hành đa kho phức tạp.")
add_bullet(doc, "Các quy định về thanh toán, điểm tích lũy và chính sách đổi trả được mô phỏng theo phạm vi đồ án, không thay thế hệ thống thương mại thực tế.")

add_heading(doc, "2.2.5. Tiêu chí đánh giá hoàn thành phạm vi", 3)
add_paragraph(
    doc,
    "Dự án được xem là hoàn thành phạm vi khi các chức năng cốt lõi của người dùng và quản trị viên hoạt động ổn định, dữ liệu được lưu "
    "đúng cấu trúc, thao tác đặt hàng không phát sinh lỗi nghiệp vụ nghiêm trọng và các thông báo quan trọng được gửi chính xác. Ngoài ra, "
    "giao diện cần thể hiện được sự nhất quán giữa các màn hình, các quy trình chính phải có thể thực hiện liên tục từ đầu đến cuối và "
    "toàn bộ nội dung báo cáo phải phản ánh đúng cách hệ thống được thiết kế, phát triển và kiểm thử.",
)

# Footer on all sections
for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not footer.text:
        add_page_number(footer)
    else:
        footer.clear()
        add_page_number(footer)

# Fix all run fonts in existing paragraphs
for para in doc.paragraphs:
    for run in para.runs:
        if run.font.name is None:
            run.font.name = "Times New Roman"

doc.save(OUTPUT)
print(OUTPUT)
